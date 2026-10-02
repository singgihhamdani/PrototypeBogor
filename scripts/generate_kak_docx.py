import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def create_revised_kak_docx(output_path):
    doc = Document()

    # Set Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("KERANGKA ACUAN KERJA (KAK) — SIJAKON KABUPATEN BOGOR TA 2026")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Dinas Pekerjaan Umum (DPU) Kabupaten Bogor")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Helper Functions for Formatting
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 81, 50) # Emerald Dark
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_meta(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(71, 85, 105)
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 81, 50)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_p(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
            rb.font.name = "Calibri"
            rb.font.size = Pt(11)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_bullet(text, level=0):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_num(text):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # --- COVER HEADER ---
    add_title("KERANGKA ACUAN KERJA (KAK) / TERM OF REFERENCE (TOR)")
    add_subtitle("PEKERJAAN: SISTEM INFORMASI JASA KONSTRUKSI KABUPATEN BOGOR")
    add_meta("SUMBER DANA: APBD KABUPATEN BOGOR TAHUN ANGGARAN 2026")
    add_meta("PEMERINTAH KABUPATEN BOGOR — DINAS PEKERJAAN UMUM DAN PENATAAN RUANG")
    
    div = doc.add_paragraph()
    div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div.paragraph_format.space_after = Pt(12)
    r_div = div.add_run("—" * 60)
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    # --- 1. LATAR BELAKANG ---
    add_heading_1("1. LATAR BELAKANG")
    add_p("Penyelenggaraan jasa konstruksi merupakan salah satu sektor strategis dalam mendukung pembangunan daerah yang berkualitas, aman, tertib, dan berkelanjutan. Pemerintah daerah memiliki kewenangan untuk melakukan pembinaan jasa konstruksi melalui penyelenggaraan sistem informasi, pengembangan sumber daya manusia, pemberdayaan badan usaha, pengawasan penyelenggaraan konstruksi, serta penyediaan data dan informasi yang akurat sebagai dasar pengambilan kebijakan.")
    add_p("Seiring meningkatnya kompleksitas pembangunan di Kabupaten Bogor (yang mencakup 40 Kecamatan), kebutuhan terhadap sistem informasi yang mampu mengintegrasikan data badan usaha jasa konstruksi (BUJK), tenaga kerja konstruksi (TKK), pelatihan, sertifikasi, pemantauan kurva S, pemetaan geospasial WebGIS, dan kegiatan pengawasan tertib konstruksi menjadi semakin krusial. Saat ini pengelolaan data masih tersebar pada berbagai media sehingga menyebabkan keterbatasan dalam penyediaan informasi yang cepat, akurat, dan terintegrasi.")
    add_p("Oleh karena itu, Pemerintah Kabupaten Bogor mengembangkan Sistem Informasi Jasa Konstruksi (SIJAKON) Kabupaten Bogor sebagai platform mandiri daerah yang mengintegrasikan seluruh basis data pembinaan, pemetaan spasial sebaran proyek, instrumen pengawasan berbasis Permen PUPR No. 1/2023, serta penyediaan fasilitas ekspor data terstandar (Microsoft Excel, CSV, PDF, dan Shapefile .SHP).")
    add_p("Pengembangan sistem informasi ini diharapkan dapat meningkatkan efektivitas pelayanan publik, transparansi, akuntabilitas, efisiensi pengelolaan data, serta mendukung proses perencanaan, monitoring, evaluasi, dan pelaporan penyelenggaraan pembinaan jasa konstruksi di Kabupaten Bogor.")

    # --- 2. MAKSUD, TUJUAN DAN SASARAN ---
    add_heading_1("2. MAKSUD, TUJUAN DAN SASARAN")
    add_heading_2("2.1 Maksud")
    add_p("Maksud dari pekerjaan ini adalah Menyusun Sistem Informasi Jasa Konstruksi Kabupaten Bogor sebagai media pengelolaan data, informasi, pelayanan, monitoring, evaluasi, dan pelaporan penyelenggaraan pembinaan jasa konstruksi secara terintegrasi.")
    
    add_heading_2("2.2 Tujuan")
    add_p("Tujuan dari pekerjaan adalah:")
    add_num("Mengidentifikasi kebutuhan sistem informasi jasa konstruksi Kabupaten Bogor.")
    add_num("Menyusun desain arsitektur sistem informasi modern berbasis web dan spasial.")
    add_num("Mengembangkan aplikasi berbasis web yang responsif, aman, dan berkinerja tinggi.")
    add_num("Mengintegrasikan pengelolaan data jasa konstruksi (BUJK, TKK, Pelatihan, Proyek, Pengawasan).")
    add_num("Meningkatkan efektivitas pembinaan dan pengawasan jasa konstruksi.")
    add_num("Menyediakan dashboard informasi dan WebGIS sebagai bahan pengambilan keputusan eksekutif.")
    add_num("Mendukung digitalisasi pelayanan publik bidang jasa konstruksi.")

    add_heading_2("2.3 Sasaran")
    add_p("Sasaran kegiatan meliputi:")
    add_bullet("Tersusunnya dokumen kebutuhan dan arsitektur sistem informasi jasa konstruksi;")
    add_bullet("Tersedianya aplikasi Sistem Informasi Jasa Konstruksi Kabupaten Bogor berbasis web;")
    add_bullet("Tersedianya basis data spasial dan tabular jasa konstruksi yang terintegrasi (PostgreSQL + PostGIS);")
    add_bullet("Meningkatnya kualitas pelayanan informasi publik dan kemudahan verifikasi legalitas;")
    add_bullet("Meningkatnya kualitas monitoring progres fisik proyek dan evaluasi kepatuhan tertib konstruksi.")

    # --- 3. NAMA PENGGUNA JASA ---
    add_heading_1("3. NAMA PENGGUNA JASA")
    add_p("Pengguna Jasa adalah: Dinas Pekerjaan Umum (DPU) Kabupaten Bogor.")

    # --- 4. DASAR HUKUM ---
    add_heading_1("4. DASAR HUKUM")
    add_p("Pelaksanaan kegiatan ini mengacu pada ketentuan peraturan perundang-undangan, antara lain:")
    add_num("Undang-Undang Nomor 2 Tahun 2017 tentang Jasa Konstruksi sebagaimana telah diubah dengan Undang-Undang Nomor 6 Tahun 2023 tentang Penetapan Perppu Cipta Kerja menjadi Undang-Undang;")
    add_num("Peraturan Pemerintah Republik Indonesia Nomor 22 Tahun 2020 tentang Peraturan Pelaksanaan UU No. 2/2017 tentang Jasa Konstruksi jo PP No. 14 Tahun 2021;")
    add_num("Peraturan Menteri Pekerjaan Umum dan Perumahan Rakyat Nomor 9 Tahun 2020 tentang Pembentukan Lembaga Pengembangan Jasa Konstruksi;")
    add_num("Peraturan Menteri Pekerjaan Umum dan Perumahan Rakyat Nomor 1 Tahun 2023 tentang Pedoman Pengawasan Penyelenggaraan Jasa Konstruksi yang Dilaksanakan Pemerintah Daerah Provinsi, Kabupaten, dan Kota;")
    add_num("Peraturan Daerah Provinsi Jawa Barat Nomor 6 Tahun 2024 tentang Pembinaan dan Pengawasan Jasa Konstruksi.")

    # --- 5. LINGKUP KEGIATAN ---
    add_heading_1("5. LINGKUP KEGIATAN")
    add_p("Lingkup kegiatan pada Penyusunan Sistem Informasi Jasa Konstruksi meliputi tahapan terstruktur sebagai berikut:")
    
    add_heading_2("A. Tahap Perencanaan & Persiapan")
    add_bullet("Kick-Off Meeting dan penyelarasan pemahaman teknis dengan DPU Kab. Bogor.")
    add_bullet("Penyusunan rencana kerja, jadwal pelaksanaan 90 hari, dan metodologi pengembangan.")
    add_bullet("Inventarisasi data eksisting, studi regulasi, dan identifikasi kebutuhan stakeholder.")

    add_heading_2("B. Tahap Perancangan & Analisis Sistem")
    add_bullet("Analisis proses bisnis tata kelola jasa konstruksi dan instrumen audit Permen PUPR 1/2023.")
    add_bullet("Desain arsitektur sistem, basis data relasional + PostGIS spasial, dan keamanan RBAC.")
    add_bullet("Desain antarmuka pengguna (UI/UX) modern, Side-by-Side Document Reviewer, dan dashboard eksekutif.")

    add_heading_2("C. Tahap Pengembangan Aplikasi")
    add_bullet("Pengembangan modul core: Autentikasi, Manajemen Pengguna, dan Audit Trail Activity Log.")
    add_bullet("Pengembangan Modul Master BUJK: Registrasi mandiri, SBU, pengalaman proyek, dan kurva S.")
    add_bullet("Pengembangan Modul WebGIS: Pemetaan sebaran proyek 40 Kecamatan dan parser Shapefile (.SHP).")
    add_bullet("Pengembangan Modul Pelatihan TKK: Pendaftaran online, seleksi peserta, dan e-Certificate ber-QR Code.")
    add_bullet("Pengembangan Modul Pengawasan: Checklist digital 3 Tertib dan fasilitas unggah instrumen SIMAK.")
    add_bullet("Pengembangan Modul Pelaporan Eksekutif: Ekspor data multi-format (Excel, CSV, PDF, SHP).")

    add_heading_2("D. Tahap Implementasi & Pengujian")
    add_bullet("Pengujian sistem secara menyeluruh: Unit Testing, Functional Testing, Security & Performance, serta UAT.")
    add_bullet("Instalasi dan deployment pada infrastruktur server produksi, konfigurasi SSL/HTTPS, dan backup otomatis.")
    add_bullet("Migrasi dan input data awal BUJK serta proyek konstruksi eksisting.")
    add_bullet("Pelatihan administrator/operator dinas dan penyusunan dokumentasi teknis (Manual Book & SOP).")

    # --- 6. LOKASI DAN SUMBER DANA ---
    add_heading_1("6. LOKASI DAN SUMBER DANA")
    add_p("Lokasi pelaksanaan pekerjaan adalah Kabupaten Bogor di lingkungan Dinas Pekerjaan Umum dan Penataan Ruang Kabupaten Bogor.")
    add_p("Sumber dana anggaran kegiatan berasal dari APBD Kabupaten Bogor Tahun Anggaran 2026.")

    # --- 7. SYARAT KUALIFIKASI PENYEDIA ---
    add_heading_1("7. PERSYARATAN KUALIFIKASI PENYEDIA")
    add_bullet("Memiliki Nomor Induk Berusaha (NIB) dengan kode KBLI 62019 : AKTIVITAS PEMROGRAMAN KOMPUTER LAINNYA;")
    add_bullet("Sertifikat Badan Usaha (SBU) Kualifikasi KECIL dengan subbidang : TELEMATIKA (1.03.05);")
    add_bullet("Status Surat Keterangan Status Wajib Pajak (KSWP) berstatus VALID melalui sistem DJP CoreTax;")
    add_bullet("Memiliki kinerja penyedia berpredikat Baik/Sangat Baik dalam 3 tahun terakhir pada SIKaP LKPP;")
    add_bullet("Memiliki pengalaman pekerjaan di bidang Jasa Konsultansi/TI sejenis minimal 1 pekerjaan dalam 3 tahun terakhir, dengan nilai tertinggi minimal 50% dari total HPS/Pagu Anggaran.")

    # --- 8. SUMBER DAYA MANUSIA (TENAGA AHLI & TENAGA PENDUKUNG) ---
    add_heading_1("8. KEBUTUHAN SUMBER DAYA MANUSIA (PERSONIL)")
    add_p("Untuk melaksanakan pekerjaan ini, Penyedia Jasa wajib menyediakan personil dengan kualifikasi sebagai berikut:")

    # Table Personil
    headers_sdm = ["No", "Posisi / Peran", "Kualifikasi Pendidikan & Pengalaman", "Tanggung Jawab Utama"]
    data_sdm = [
        ["1", "Project Manager (Ahli Informatika)", "S1 Sarjana Informatika\nPengalaman min. 3 Tahun", "Memimpin manajemen proyek, quality control, koordinasi tim, dan komunikasi berkala dengan PPK/DPU."],
        ["2", "Web Developer", "S1 Sarjana Informatika / TI / DKV\nPengalaman min. 3 Tahun", "Merancang arsitektur backend/frontend, REST API, konfigurasi server/deployment, RBAC, dan UI/UX responsif."],
        ["3", "GIS Specialist", "S1 Sarjana Geodesi / PWK\nPengalaman min. 3 Tahun", "Mengembangkan modul WebGIS, layer peta tematik 40 kecamatan, integrasi Shapefile (.SHP), dan analisis spasial."],
        ["4", "Admin Kantor", "SMK / SMA Sederajat\nPengalaman min. 1 Tahun", "Administrasi persuratan, input data awal, dokumentasi laporan kemajuan, dan penyiapan berkas serah terima."]
    ]
    
    tbl_sdm = doc.add_table(rows=len(data_sdm) + 1, cols=4)
    tbl_sdm.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sdm, color="CBD5E1", sz="4")
    
    for idx, h in enumerate(headers_sdm):
        cell = tbl_sdm.cell(0, idx)
        set_cell_background(cell, "0F5132")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row in enumerate(data_sdm):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = tbl_sdm.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx in (0, 1):
                r.font.bold = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(6)

    # --- 9. DATA PENUNJANG & STANDAR TEKNIS ---
    add_heading_1("9. DATA PENUNJANG DAN STANDAR TEKNIS YANG DIGUNAKAN")
    
    add_heading_2("9.1 Arsitektur & Bahasa Pemrograman")
    add_bullet("Sistem dikembangkan berbasis web (web-based application) yang dapat diakses melalui web browser modern tanpa memerlukan instalasi aplikasi khusus.")
    add_bullet("Pengembangan sistem menggunakan Fullstack Modern TypeScript Architecture (Next.js 15 / React 19 untuk Frontend dan NestJS / Node.js LTS untuk Backend) atau PHP versi modern (PHP 8.2+ / Laravel 11) yang mendukung arsitektur modular, Server-Side Rendering (SSR), performa tinggi, dan kemudahan pemeliharaan.")

    add_heading_2("9.2 Basis Data & Sistem Informasi Geografis (SIG)")
    add_bullet("Sistem menggunakan Database Management System PostgreSQL 16 dengan ekstensi spasial PostGIS 3.4 (EPSG:4326) untuk menyimpan data tabular dan geometri spasial (Point, MultiPolygon).")
    add_bullet("WebGIS terintegrasi penuh untuk menampilkan sebaran lokasi paket proyek konstruksi dan kantor BUJK di 40 Kecamatan Kabupaten Bogor.")
    add_bullet("Sistem mendukung fitur import dan export data spasial dalam format Shapefile (.SHP zipped) dan GeoJSON untuk interoperabilitas dengan GIS DPU / Bappedalitbang.")

    add_heading_2("9.3 Modul Pengawasan Tertib Jasa Konstruksi (Permen PUPR 1/2023)")
    add_bullet("Menyediakan lembar checklist audit digital untuk 3 lingkup: (1) Tertib Usaha, (2) Tertib Penyelenggaraan, dan (3) Tertib Pemanfaatan Produk Konstruksi.")
    add_bullet("Fasilitas unggah berkas excel instrumen SIMAK dan kalkulasi otomatis skor kepatuhan pengawasan.")

    add_heading_2("9.4 Modul Pelatihan TKK & e-Certificate ber-QR Code")
    add_bullet("Registrasi online mandiri calon peserta pelatihan TKK dengan validasi NIK dan berkas persyaratan.")
    add_bullet("Generator e-Certificate otomatis (PDF) dengan penomoran dinas resmi dan QR Code dinamis untuk validasi keabsahan secara online oleh publik.")

    add_heading_2("9.5 Antarmuka Pengguna & Inovasi Interaksi (UI/UX)")
    add_bullet("Desain responsif menyesuaikan perangkat desktop, laptop, tablet, dan ponsel pintar.")
    add_bullet("Fitur Side-by-Side Document Reviewer (In-App PDF Viewer berdampingan dengan form verifikasi) untuk mempercepat proses persetujuan berkas pendaftaran BUJK.")
    add_bullet("Multi-Step Stepper Wizard pada formulir pendaftaran BUJK dan pelatihan.")

    add_heading_2("9.6 Keamanan Sistem & Audit Trail")
    add_bullet("Manajemen hak akses berbasis peran (Role-Based Access Control / RBAC) dengan 3 tingkatan utama: Visitor, Operator, dan Super Admin.")
    add_bullet("Otentikasi aman menggunakan enkripsi password standar industri (Argon2id / Bcrypt) dan token JWT.")
    add_bullet("Pencatatan riwayat aktivitas pengguna (Audit Trail Log) untuk seluruh aksi penambahan, pengubahan, dan penghapusan data.")
    add_bullet("Koneksi jaringan wajib menggunakan protokol aman HTTPS / SSL.")

    # --- 10. KELUARAN / OUTPUT ---
    add_heading_1("10. KELUARAN / OUTPUT PEKERJAAN")
    add_p("Keluaran dari pekerjaan ini meliputi:")
    add_bullet("Aplikasi Sistem Informasi Jasa Konstruksi berbasis web yang responsif, aman, dan berkinerja tinggi;")
    add_bullet("Dashboard eksekutif, peta tematik WebGIS 40 Kecamatan, dan statistik interaktif;")
    add_bullet("Modul Master BUJK, SBU, Portofolio Pengalaman, dan Kurva S Progres Proyek;")
    add_bullet("Modul Pengawasan Tertib Konstruksi (Permen PUPR 1/2023) dan unggah SIMAK;")
    add_bullet("Modul Pelatihan TKK, e-Certificate digital (PDF), dan QR Code verifikasi publik;")
    add_bullet("Modul Pelaporan Eksekutif dan ekspor data multi-format (Microsoft Excel, CSV, PDF, dan Shapefile .SHP);")
    add_bullet("Basis data terstruktur PostgreSQL + PostGIS 3.4 beserta dokumentasi skema relasi data;")
    add_bullet("Dokumentasi teknis sistem, Buku Panduan Pengguna (Manual Book), dan SOP pengoperasian;")
    add_bullet("Source code lengkap aplikasi sesuai ketentuan kontrak dan Berita Acara UAT.")

    # --- 11. JADWAL PELAKSANAAN (90 HARI) ---
    add_heading_1("11. JADWAL PELAKSANAAN PEKERJAAN (90 HARI KALENDER)")
    add_p("Pelaksanaan pekerjaan direncanakan selama 90 (sembilan puluh) hari kalender dengan matriks tahapan sebagai berikut:")

    headers_jdw = ["No", "Tahapan / Kegiatan Utama", "H 1-15", "H 16-30", "H 31-45", "H 46-60", "H 61-75", "H 76-90"]
    data_jdw = [
        ["1", "Mobilisasi, Kick-off & Penyusunan Metodologi", "■", "", "", "", "", ""],
        ["2", "Inventarisasi Data & Analisis Kebutuhan Sistem", "■", "■", "", "", "", ""],
        ["3", "Perancangan Arsitektur, Basis Data & UI/UX", "", "■", "■", "", "", ""],
        ["4", "Pengembangan Modul Core, RBAC & Master BUJK", "", "■", "■", "", "", ""],
        ["5", "Pengembangan WebGIS 40 Kec. & Format SHP", "", "", "■", "■", "", ""],
        ["6", "Pengembangan Pelatihan TKK & e-Certificate QR", "", "", "■", "■", "", ""],
        ["7", "Pengembangan Modul Pengawasan Permen 1/2023", "", "", "", "■", "■", ""],
        ["8", "Pengembangan Modul Pelaporan & Ekspor Data", "", "", "", "■", "■", ""],
        ["9", "Testing (Functional, Security) & UAT Stakeholder", "", "", "", "", "■", "■"],
        ["10", "Migrasi Data, Pelatihan Admin & Serah Terima", "", "", "", "", "", "■"]
    ]
    
    tbl_jdw = doc.add_table(rows=len(data_jdw) + 1, cols=8)
    tbl_jdw.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_jdw, color="CBD5E1", sz="4")
    
    for idx, h in enumerate(headers_jdw):
        cell = tbl_jdw.cell(0, idx)
        set_cell_background(cell, "0F5132")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p = cell.paragraphs[0]
        if idx >= 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row in enumerate(data_jdw):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = tbl_jdw.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            if c_idx >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 0:
                r.font.bold = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(6)

    # --- 12. ALIH PENGETAHUAN & PENUTUP ---
    add_heading_1("12. ALIH PENGETAHUAN")
    add_p("Penyedia Jasa wajib menyelenggarakan kegiatan alih pengetahuan melalui bimbingan teknis, pelatihan administrator/operator, dan pendampingan pengoperasian sistem kepada aparatur Dinas Pekerjaan Umum dan Penataan Ruang Kabupaten Bogor guna menjamin keberlanjutan operasional sistem secara mandiri.")

    # Sign-off Box
    sp_sign = doc.add_paragraph()
    sp_sign.paragraph_format.space_before = Pt(20)
    
    tbl_sign = doc.add_table(rows=1, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sign, color="FFFFFF", sz="0", val="none")
    
    # Left Empty
    c_left = tbl_sign.cell(0, 0)
    # Right Sign
    c_right = tbl_sign.cell(0, 1)
    
    p_s = c_right.paragraphs[0]
    p_s.paragraph_format.line_spacing = 1.15
    r_s1 = p_s.add_run("Cibinong, Agustus 2026\nPejabat Pembuat Komitmen (PPK)\nDPU Kabupaten Bogor\n\n\n\n\n")
    r_s1.font.name = "Calibri"
    r_s1.font.size = Pt(11)
    
    r_s2 = p_s.add_run("Bang Fauzy Tea\n")
    r_s2.font.name = "Calibri"
    r_s2.font.size = Pt(11)
    r_s2.font.bold = True
    r_s2.font.underline = True
    
    r_s3 = p_s.add_run("NIP. ----------------------------")
    r_s3.font.name = "Calibri"
    r_s3.font.size = Pt(10)

    doc.save(output_path)
    print(f"KAK Revisi berhasil disimpan di: {output_path}")

if __name__ == "__main__":
    out_docx = r"u:\Project\ciptabintar\KAK APLIKASI JAKON BOGOR 2026_REVISI.docx"
    create_revised_kak_docx(out_docx)
