import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def generate_planning_kak_docx(input_md_path, output_docx_path):
    doc = Document()

    # Set Margins (1 inch / 2.54cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("KAK JASA KONSULTANSI PERENCANAAN — SIJAKON KAB. BOGOR TA 2026")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Dinas Pekerjaan Umum dan Penataan Ruang (DPUPR) Kabupaten Bogor | Dokumen Pengadaan Perencanaan")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    with open(input_md_path, 'r', encoding='utf-8') as f:
        md_content = f.read()

    lines = md_content.split('\n')
    
    # Title Section
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_t = title_p.add_run("KERANGKA ACUAN KERJA (KAK) / TERM OF REFERENCE (TOR)")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(15)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(15, 81, 50) # Emerald Dark #0F5132

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(2)
    run_sub = sub_p.add_run("JASA KONSULTANSI PERENCANAAN:\nPENYUSUNAN RANCANGAN DAN DED SISTEM INFORMASI JASA KONSTRUKSI (SIJAKON)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(30, 41, 59)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(12)
    run_meta = meta_p.add_run("SUMBER DANA: APBD KABUPATEN BOGOR TAHUN ANGGARAN 2026\nPEMERINTAH KABUPATEN BOGOR — DINAS PEKERJAAN UMUM DAN PENATAAN RUANG")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(10.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(71, 85, 105)

    div_p = doc.add_paragraph()
    div_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div_p.paragraph_format.space_after = Pt(14)
    r_div = div_p.add_run("—" * 60)
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip headers already added
        if stripped.startswith('# KERANGKA ACUAN KERJA') or stripped.startswith('## JASA KONSULTANSI') or stripped.startswith('### TAHUN ANGGARAN') or stripped == '---':
            i += 1
            continue

        # Heading 2 (##)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(16)
            hp.paragraph_format.space_after = Pt(6)
            hp.paragraph_format.keep_with_next = True
            hrun = hp.add_run(h_text)
            hrun.font.name = "Arial"
            hrun.font.size = Pt(13)
            hrun.font.bold = True
            hrun.font.color.rgb = RGBColor(15, 81, 50)
            i += 1
            continue

        # Heading 3 (###)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(12)
            hp.paragraph_format.space_after = Pt(4)
            hp.paragraph_format.keep_with_next = True
            hrun = hp.add_run(h_text)
            hrun.font.name = "Arial"
            hrun.font.size = Pt(11.5)
            hrun.font.bold = True
            hrun.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Table handler
        if stripped.startswith('|') and '|' in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                data_rows = []
                for row_line in table_lines[2:]:
                    cols = [c.strip() for c in row_line.split('|')[1:-1]]
                    if len(cols) == len(headers):
                        data_rows.append(cols)
                
                doc_table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(doc_table, color="CBD5E1", sz="4")
                
                # Header row
                hdr_cells = doc_table.rows[0].cells
                for idx, heading in enumerate(headers):
                    cell = hdr_cells[idx]
                    set_cell_background(cell, "0F5132")
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    r = p.add_run(heading.replace('**', ''))
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                
                # Data rows
                for r_idx, row_data in enumerate(data_rows):
                    row_cells = doc_table.rows[r_idx + 1].cells
                    bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    for c_idx, cell_value in enumerate(row_data):
                        cell = row_cells[c_idx]
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1.15
                        
                        clean_text = cell_value.replace('<br>', '\n')
                        clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', clean_text)
                        
                        parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
                        for part in parts:
                            if part.startswith('**') and part.endswith('**'):
                                r = p.add_run(part[2:-2])
                                r.font.bold = True
                            else:
                                r = p.add_run(part)
                            r.font.name = "Calibri"
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = RGBColor(30, 41, 59)
                
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(4)
                sp.paragraph_format.space_after = Pt(6)
            continue

        # Bullet lists (* or -)
        if stripped.startswith('* ') or stripped.startswith('- '):
            list_text = stripped[2:].strip()
            lp = doc.add_paragraph(style='List Bullet')
            lp.paragraph_format.space_before = Pt(0)
            lp.paragraph_format.space_after = Pt(3)
            lp.paragraph_format.line_spacing = 1.15
            
            clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', list_text)
            parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = lp.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = lp.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Numbered list (1. 2. etc)
        match_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if match_num:
            num_text = match_num.group(2)
            lp = doc.add_paragraph(style='List Number')
            lp.paragraph_format.space_before = Pt(0)
            lp.paragraph_format.space_after = Pt(3)
            lp.paragraph_format.line_spacing = 1.15
            
            clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', num_text)
            parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = lp.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = lp.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Regular Paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', stripped)
            parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.font.italic = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

        i += 1

    doc.save(output_docx_path)
    print(f"KAK Perencanaan berhasil dibuat di: {output_docx_path}")

if __name__ == "__main__":
    md_in = r"u:\Project\ciptabintar\KAK_PERENCANAAN_SISTEM_INFORMASI_JAKON_BOGOR_2026.md"
    docx_out = r"u:\Project\ciptabintar\KAK_PERENCANAAN_SISTEM_INFORMASI_JAKON_BOGOR_2026.docx"
    doc_out = r"u:\Project\ciptabintar\KAK_PERENCANAAN_SISTEM_INFORMASI_JAKON_BOGOR_2026.doc"
    
    generate_planning_kak_docx(md_in, docx_out)
    
    import shutil
    try:
        shutil.copyfile(docx_out, doc_out)
        print(f"File .doc berhasil disalin ke: {doc_out}")
    except PermissionError:
        print("Info: File .doc sedang dibuka di Word.")
