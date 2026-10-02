import os
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, title, text_items):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Left border emerald
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="0F5132"/>'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r_t = p.add_run(title)
    r_t.font.name = "Arial"
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 81, 50)
    
    for item in text_items:
        ip = cell.add_paragraph()
        ip.paragraph_format.space_before = Pt(0)
        ip.paragraph_format.space_after = Pt(3)
        ir = ip.add_run(item)
        ir.font.name = "Calibri"
        ir.font.size = Pt(10)
        ir.font.color.rgb = RGBColor(51, 65, 85)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(2)
    sp.paragraph_format.space_after = Pt(4)

def build_user_hierarchy_table(doc):
    headers = ["Tingkat Hak Akses", "Kategori Pengguna", "Hak Akses & Fitur Utama"]
    data = [
        ["1. Visitor / Publik", "Masyarakat, Umum, Calon Peserta", "• Akses Peta Sebaran Proyek WebGIS 40 Kecamatan\n• Katalog Regulasi & Publikasi Berita Jakon\n• Pendaftaran Akun BUJK & Pendaftaran Bimtek TKK\n• Pemindaian & Validasi QR Code e-Certificate"],
        ["2. Operator", "Operator BUJK, PPK / OPD, Asesor Pengawas", "• Operator BUJK: Input SBU, Pengalaman, SIMAK, Progres Fisik\n• Operator OPD/PPK: Input Data Paket Pekerjaan Konstruksi APBD\n• Asesor Pengawas: Pengisian Checklist & Skor Permen 1/2023"],
        ["3. Super Admin", "Admin Dinas PUPR Kab. Bogor", "• Verifikasi & Approval Pendaftaran Akun & Berkas Legalitas\n• Penjadwalan Pengawasan Tahunan & Master Bimtek Pelatihan\n• Manajemen Matriks RBAC & Pemantau Audit Trail Log\n• Ekspor Data Multi-Format (Shapefile .SHP, Excel, CSV, PDF)"]
    ]
    
    table = doc.add_table(rows=len(data) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CBD5E1", sz="4")
    
    # Header
    for idx, h in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_background(cell, "0F5132")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row in enumerate(data):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 0:
                r.font.bold = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(6)

def build_architecture_table(doc):
    headers = ["Lapisan Arsitektur", "Teknologi Utama", "Fungsi & Peran dalam Sistem"]
    data = [
        ["Frontend (Client Layer)", "Next.js 15 (App Router, RSC) + React 19 + Tailwind CSS + Shadcn UI", "Antarmuka modern responsif, Server Components untuk kecepatan muat, Side-by-Side Reviewer, dan Stepper Wizard."],
        ["WebGIS Engine", "MapLibre GL JS / Leaflet + shpjs + @turf/turf", "Pemetaan spasial GPU-accelerated sebaran proyek 40 kecamatan, import/export Shapefile (.SHP) dan GeoJSON."],
        ["Backend (Business Layer)", "NestJS / Hono.js (Node.js LTS, TypeScript)", "Arsitektur modular, Controller-Service-Repository, validasi skema Zod, otentikasi JWT, dan RBAC granular."],
        ["Database & Spatial Layer", "PostgreSQL 16 + PostGIS 3.4 Extension", "Penyimpanan data relasional dan tipe data geometri spasial (Point, Polygon, EPSG:4326) dengan indeks GiST cepat."],
        ["Cache & Background Queue", "Redis 7 + BullMQ", "Antrian background jobs untuk generator PDF e-Certificate massal, ekspor file spasial SHP, dan automated expiry alert."],
        ["File & Document Engine", "Object Storage (MinIO/S3) + Puppeteer + ExcelJS", "Penyimpanan PDF legalitas terenkripsi, pembuatan sertifikat ber-QR code, dan ekspor spreadsheet Excel/CSV."]
    ]
    
    table = doc.add_table(rows=len(data) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CBD5E1", sz="4")
    
    for idx, h in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_background(cell, "0F5132")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row in enumerate(data):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 0:
                r.font.bold = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(6)

def build_timeline_table(doc):
    headers = ["Fase Pelaksanaan", "Alokasi Waktu", "Rincian Aktivitas & Output Deliverable"]
    data = [
        ["Fase 1: Persiapan & Analisis", "Hari 1 - 15 (2 Minggu)", "• Kick-off meeting teknis dengan DPU Kab. Bogor & inventarisasi data awal.\n• Analisis proses bisnis Permen PUPR 1/2023, data spasial 40 Kecamatan, dan skema database."],
        ["Fase 2: Perancangan & Core System", "Hari 16 - 35 (3 Minggu)", "• Perancangan UI/UX Design System & Prototipe Interaktif Dashboard.\n• Pengembangan Autentikasi, Granular RBAC, Audit Trail, dan Pendaftaran BUJK/TKK.\n• Pembangunan Master BUJK, SBU, Portofolio Pengalaman, dan Kurva S."],
        ["Fase 3: WebGIS & Pelatihan TKK", "Hari 36 - 55 (3 Minggu)", "• Integrasi WebGIS sebaran proyek & kantor BUJK di 40 Kecamatan.\n• Pengembangan modul parser & converter format spasial Shapefile (.SHP) dan GeoJSON.\n• Pembangunan modul Pelatihan TKK, seleksi pendaftar, dan e-Certificate ber-QR Code."],
        ["Fase 4: Pengawasan & Pelaporan", "Hari 56 - 75 (3 Minggu)", "• Implementasi checklist audit digital (Tertib Usaha, Penyelenggaraan, Pemanfaatan).\n• Fasilitas upload SIMAK, scoring otomatis, dan modul Pelaporan Eksekutif.\n• Pembuatan generator rekapitulasi data periodik dan import data massal Excel."],
        ["Fase 5: Testing, Training & Go-Live", "Hari 76 - 90 (2 Minggu)", "• Pengujian menyeluruh (Functional Testing, Security & Performance, UAT).\n• Migrasi data awal, konfigurasi server produksi, penyusunan Manual Book & SOP.\n• Pelatihan administrator dinas dan serah terima aplikasi 100%."]
    ]
    
    table = doc.add_table(rows=len(data) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="CBD5E1", sz="4")
    
    for idx, h in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_background(cell, "0F5132")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row in enumerate(data):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(30, 41, 59)
            if c_idx == 0:
                r.font.bold = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(4)
    sp.paragraph_format.space_after = Pt(6)

def generate_full_prd_docx(input_md_path, output_docx_path):
    doc = Document()

    # Set Margins (1 inch / 2.54cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SIJAKON BOGOR TA 2026 — Product Requirement Document (PRD)")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Dinas Pekerjaan Umum (DPU) Kabupaten Bogor | Dokumen Perencanaan Teknis")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    with open(input_md_path, 'r', encoding='utf-8') as f:
        md_content = f.read()

    lines = md_content.split('\n')
    
    # Title Section
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_t = title_p.add_run("Product Requirement Document (PRD) Rekomendasi")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(20)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(15, 81, 50) # Emerald Dark #0F5132

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Sistem Informasi Pembinaan & Pengawasan Jasa Konstruksi (SIJAKON) Kabupaten Bogor TA 2026")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(14)
    r_div = div_p.add_run("—" * 55)
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    i = 0
    in_code_block = False
    code_type = ""
    code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip headers already added
        if stripped.startswith('# Product Requirement Document') or stripped.startswith('# Sistem Informasi Jasa Konstruksi') or stripped == '---':
            i += 1
            continue

        # Code block handler (Mermaid, Text, Schema)
        if stripped.startswith('```'):
            if in_code_block:
                # End of code block - render appropriate clean Word representation!
                in_code_block = False
                
                if "graph TD" in '\n'.join(code_lines):
                    # Section 2 User Hierarchy
                    build_user_hierarchy_table(doc)
                elif "FRONTEND (CLIENT LAYER)" in '\n'.join(code_lines):
                    # Section 3 Architecture Blueprint
                    build_architecture_table(doc)
                elif "FASE 1: PERSIAPAN" in '\n'.join(code_lines):
                    # Section 8 Timeline
                    build_timeline_table(doc)
                else:
                    # Generic clean code container
                    code_text = '\n'.join(code_lines)
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = tbl.cell(0, 0)
                    set_cell_background(cell, "F8FAFC")
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                    
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    crun = cp.add_run(code_text)
                    crun.font.name = 'Consolas'
                    crun.font.size = Pt(8.5)
                    crun.font.color.rgb = RGBColor(30, 41, 59)
                    
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_before = Pt(4)
                    sp.paragraph_format.space_after = Pt(6)
                
                code_lines = []
            else:
                in_code_block = True
                code_type = stripped[3:].strip()
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Heading 2 (##)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(16)
            hp.paragraph_format.space_after = Pt(6)
            hp.paragraph_format.keep_with_next = True
            hrun = hp.add_run(h_text)
            hrun.font.name = "Arial"
            hrun.font.size = Pt(14)
            hrun.font.bold = True
            hrun.font.color.rgb = RGBColor(15, 81, 50) # Emerald Dark
            i += 1
            continue

        # Heading 3 (###)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = Pt(12)
            hp.paragraph_format.space_after = Pt(4)
            hp.paragraph_format.keep_with_next = True
            hrun = hp.add_run(h_text)
            hrun.font.name = "Arial"
            hrun.font.size = Pt(12)
            hrun.font.bold = True
            hrun.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
            i += 1
            continue

        # Table handler
        if stripped.startswith('|') and '|' in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                data_rows = []
                for row_line in table_lines[2:]:
                    cols = [c.strip() for c in row_line.split('|')[1:-1]]
                    if len(cols) == len(headers):
                        data_rows.append(cols)
                
                doc_table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(doc_table, color="CBD5E1", sz="4")
                
                # Header row
                hdr_cells = doc_table.rows[0].cells
                for idx, heading in enumerate(headers):
                    cell = hdr_cells[idx]
                    set_cell_background(cell, "0F5132")
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    r = p.add_run(heading.replace('**', ''))
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                
                # Data rows
                for r_idx, row_data in enumerate(data_rows):
                    row_cells = doc_table.rows[r_idx + 1].cells
                    bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                    for c_idx, cell_value in enumerate(row_data):
                        cell = row_cells[c_idx]
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1.15
                        
                        clean_text = cell_value.replace('<br>', '\n')
                        clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', clean_text)
                        
                        parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
                        for part in parts:
                            if part.startswith('**') and part.endswith('**'):
                                r = p.add_run(part[2:-2])
                                r.font.bold = True
                            else:
                                r = p.add_run(part)
                            r.font.name = "Calibri"
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = RGBColor(30, 41, 59)
                
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(4)
                sp.paragraph_format.space_after = Pt(6)
            continue

        # Bullet lists (* or -)
        if stripped.startswith('* ') or stripped.startswith('- '):
            list_text = stripped[2:].strip()
            lp = doc.add_paragraph(style='List Bullet')
            lp.paragraph_format.space_before = Pt(0)
            lp.paragraph_format.space_after = Pt(3)
            lp.paragraph_format.line_spacing = 1.15
            
            clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', list_text)
            parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = lp.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = lp.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Numbered list (1. 2. etc)
        match_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if match_num:
            num_text = match_num.group(2)
            lp = doc.add_paragraph(style='List Number')
            lp.paragraph_format.space_before = Pt(0)
            lp.paragraph_format.space_after = Pt(3)
            lp.paragraph_format.line_spacing = 1.15
            
            clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', num_text)
            parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = lp.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = lp.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Regular Paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            clean_text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', stripped)
            parts = re.split(r'(\*\*[^*]+\*\*)', clean_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.font.italic = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(30, 41, 59)

        i += 1

    doc.save(output_docx_path)
    print(f"File berhasil di-render di: {output_docx_path}")

if __name__ == "__main__":
    md_in = r"u:\Project\ciptabintar\PRD_REKOMENDASI_JAKON_BOGOR_2026.md"
    docx_out = r"u:\Project\ciptabintar\PRD_REKOMENDASI_JAKON_BOGOR_2026.docx"
    doc_out = r"u:\Project\ciptabintar\PRD_REKOMENDASI_JAKON_BOGOR_2026.doc"
    
    generate_full_prd_docx(md_in, docx_out)
    
    try:
        shutil.copyfile(docx_out, doc_out)
        print(f"File .doc berhasil disalin ke: {doc_out}")
    except PermissionError:
        print("Info: File .doc sedang dibuka di Word. File .docx berhasil diperbarui dengan sempurna.")
